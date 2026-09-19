# -*- coding: utf-8 -*-
"""Genera Hallazgos_Analisis_Sistema.docx a partir del análisis de ingeniería inversa."""
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn


OUT = Path(r"D:\UPN\CICLO 10\Evo conf software\RRHH\RRHH\Hallazgos_Analisis_Sistema.docx")


def set_cell_shading(cell, fill_hex: str):
    """Aplica fondo a celda de tabla."""
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tc.append(shd)


def style_header_row(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    style_header_row(table.rows[0])
    for r_idx, row_data in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            cells[c_idx].text = str(val)
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    return p


def main():
    doc = Document()

    # Estilos base
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Portada / título
    title = doc.add_heading("Hallazgos de Análisis Funcional y Técnico", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Sistema de Recursos Humanos (RecursosHumanos / GRLL)\n"
        "Informe de ingeniería inversa — análisis estático de binarios, "
        "configuración y metadatos"
    )
    r.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        "Proyecto: RRHH | Artefacto principal: RecursosHumanos.exe | "
        "Base de datos referenciada: CMI (SQL Server)\n"
        "Fecha del informe: 8 de septiembre de 2026"
    )
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ========== RESUMEN EJECUTIVO ==========
    doc.add_heading("Resumen ejecutivo", level=1)
    resumen = [
        "El repositorio analizado es principalmente un despliegue compilado (.NET / VB.NET Windows Forms), no un árbol de código fuente completo.",
        "La aplicación cliente RecursosHumanos.exe gestiona escalafón, asistencia, permisos, fichas laborales/sociales y reportes sobre SQL Server (catálogo CMI).",
        "El acceso al sistema autentica al usuario como login de SQL Server (User Id/Password del formulario), no contra una tabla de aplicación en el click de Aceptar.",
        "La capa compartida GRLL.dll concentra conexión (clsConexion), entidades/mappers LINQ y cientos de contratos a procedimientos almacenados spRRHH_*/spSEG_*.",
        "ClassRRHH.dll aporta lógica de dominio RRHH (AsistenciaAccess, FichaSocialClass, etc.); Seguridad.dll solo expone cifrado TripleDES+MD5.",
        "UI basada en DevExpress v11.2 (y restos v7.3/v9.1) y Microsoft ReportViewer; arquitectura monolítica de escritorio sin API HTTP.",
        "No hay README ni scripts DDL del esquema CMI en el repositorio; tablas/PK/FK/triggers no pudieron verificarse directamente.",
        "Credenciales SQL (usuario sa) aparecen en texto plano en RecursosHumanos.exe.config — riesgo de seguridad confirmado.",
        "Punto único de falla crítico: clsConexion + SQL Server; si falla la autenticación SQL o la BD, el sistema no arranca funcionalmente.",
        "Limitaciones típicas de legado: deuda de versiones de UI, autenticación acoplada a SQL, ausencia de pruebas automatizadas y documentación funcional.",
    ]
    for line in resumen:
        p = doc.add_paragraph(line, style="List Number")
        for run in p.runs:
            run.font.size = Pt(10)

    # ========== METODOLOGÍA ==========
    doc.add_heading("Metodología", level=1)
    add_para(
        doc,
        "Alcance: análisis estático del directorio del proyecto "
        r"D:\UPN\CICLO 10\Evo conf software\RRHH\RRHH "
        "(y subcarpetas Reportes/, temp/, _auth_analysis/). "
        "No se ejecutó la aplicación contra un servidor SQL vivo ni se "
        "inspeccionó el esquema físico de la base CMI en este informe.",
    )
    add_para(doc, "Fuentes revisadas:", bold=True)
    fuentes = [
        "Binarios: RecursosHumanos.exe, GRLL.dll, ClassRRHH.dll, Seguridad.dll, ClsCtrls.dll, GRLL.Common.Rutinas.dll (+ copias en Reportes/ y temp/).",
        "Configuración: RecursosHumanos.exe.config (connectionStrings, applicationSettings, userSettings).",
        "Metadatos/PDB: rutas de origen en _auth_analysis/pdb_sources.txt (p.ej. FrmLogin.vb, ucLogin.vb, clsConexion.vb).",
        "Inventarios de tipos: _auth_analysis/exe_classes.txt, grll_classes.txt, classrrhh_classes.txt, seguridad_classes.txt.",
        "Cadenas extraídas: strings_RecursosHumanos.txt, strings_GRLL.txt.",
        "Descompilación parcial previa: dec_GRLL_Common_ucLogin.cs, dec_GRLL_Common_clsConexion.cs, dec_GRLL_Common_clsLogin.cs, dec_Seguridad.cs.",
        "Documentación XML de ensamblado: RecursosHumanos.xml, ClassRRHH.xml, GRLL.xml, Seguridad.xml (en su mayoría stubs de diseñador / recursos).",
        "Búsqueda de README/SQL/DDL: no se encontró documentación de usuario ni scripts de esquema.",
    ]
    for f in fuentes:
        doc.add_paragraph(f, style="List Bullet")

    add_para(
        doc,
        "Limitación metodológica: las afirmaciones sobre tablas físicas, "
        "constraints y triggers se marcan como no determinadas cuando solo "
        "existen indicios vía nombres de entidades/stored procedures en el "
        "ensamblado, sin script DDL ni acceso a sys.tables.",
        bold=False,
    )

    # ========== 1. FUNCIONALIDADES ==========
    doc.add_heading("1. Funcionalidades principales", level=1)
    add_para(
        doc,
        "Identificadas por nombres de formularios (Frm*), cadenas de UI, "
        "TableAdapters/DataSet1 y procedimientos spRRHH_* referenciados en "
        "RecursosHumanos.exe / GRLL.dll / RecursosHumanos.xml.",
    )

    doc.add_heading("1.1 Acceso y seguridad de aplicación", level=2)
    add_table(
        doc,
        ["Funcionalidad", "Evidencia de identificación", "Descripción observada"],
        [
            [
                "Login / Acceso al Sistema",
                "RecursosHumanos.FrmLogin; título UI 'Acceso al Sistema'; GRLL.Common.ucLogin",
                "Formulario de usuario/contraseña; botón Aceptar abre SqlConnection vía clsConexion.",
            ],
            [
                "Límite de intentos de login",
                "Cadena 'Nº intentos = 3' y contador estático en ucLogin.btnAceptar_Click",
                "Tras 3 fallos dispara onSalirIntentos.",
            ],
            [
                "Verificación de acceso de app (post-login)",
                "Cadenas 'No tiene Acceso al Sistema'; tipos GRLL_Usuario_Acceso / GRLL_Usuario_AccesoResult",
                "Indicio de chequeo de permisos de aplicación después de conectar; detalle de flujo en FrmLogin no descompilado completamente en este paquete.",
            ],
            [
                "Cifrado auxiliar",
                "Seguridad.Seguridad.EncryptString / DecryptString",
                "TripleDES + MD5 para clave; no es el mecanismo principal del login SQL.",
            ],
        ],
    )

    doc.add_heading("1.2 Escalafón / ficha de personal", level=2)
    add_table(
        doc,
        ["Funcionalidad", "Formularios / clases", "Indicios de datos"],
        [
            [
                "Nueva ficha / datos personales",
                "FrmNuevaFicha, FrmMantEmpleado_Basico",
                "spRRHH_Ficha_DatosPersonales",
            ],
            [
                "Periodo laboral, rotaciones, resoluciones",
                "FrmPeriodoLaboral, FrmRotaciones, FrmResoluciones",
                "spRRHH_GetPeriodoLaboralByEmpleado, spRRHH_ListarRotaciones, spRRHH_ListarResolucion",
            ],
            [
                "Estudios, títulos, capacitaciones",
                "FrmEstudiosRealizados, FrmTitulos, FrmCapacitaciones, FrmMantTipoEstudio, FrmMantTipoInst",
                "spRRHH_Ficha_Estudio_*, spRRHH_buscarEstudios",
            ],
            [
                "Experiencia laboral",
                "FrmExpLaboral",
                "spRRHH_Ficha_ExpLab, spRRHH_BuscarExpLab",
            ],
            [
                "Familiares / ficha social / vivienda",
                "FrmFamiliar, FrmFichaSocial, ClassRRHH.FichaSocialClass / FamiliaresClass",
                "spRRHH_Ficha_Det_Familiar*, spRRHH_Ficha_Vivienda",
            ],
            [
                "Documentos / archivos",
                "FrmDocumentos, FrmListaArchivos, FrmMantDocFile",
                "spRRHH_busca_archivos; mensajes 'Debe Seleccionar un Documento'",
            ],
            [
                "Disciplina",
                "FrmDisciplina",
                "spRRHH_busca_Disciplina",
            ],
            [
                "Relación orgánica / personal por área",
                "FrmListaOrganica",
                "UI 'UNIDADES ORGANICAS' / 'PERSONAL DE:'",
            ],
            [
                "Maestros (estado civil, sangre, nacionalidad, régimen, docs, condición laboral, cargos)",
                "FrmMantEstCivil, FrmMantGrupoSanguineo, FrmMantNacionalidades, FrmMantRegimen, FrmMantTipoDocIdent, FrmMantTipoTrab, FrmMantCargos",
                "Nombres de formularios y mensajes de validación en strings del exe",
            ],
            [
                "Asegurados / derechohabientes",
                "FrmRegistroAsegurado, FrmBuscarDerechohabiente, FrmBuscarPersona",
                "Mensajes 'Primero debe Seleccionar un Asegurado'",
            ],
        ],
    )

    doc.add_heading("1.3 Asistencia, horarios y permisos", level=2)
    add_table(
        doc,
        ["Funcionalidad", "Formularios", "SP / clases asociadas"],
        [
            [
                "Horarios y asignación",
                "FrmHorarios, FrmMantHorarios, FrmAsignarHorario, FrmAsignarHorario_Lector, FrmHorarioEsp, FrmConsultaHorarioEsp, FrmLocalAsignado",
                "Entidades Horario/HorarioTemporal; sp_Listar_Empleado_Horario, sp_Listar_Empleado_Local",
            ],
            [
                "Marcación / asistencia",
                "FrmAsistenciaManual, FrmConsultaAsistencias, FrmConsultaAsistDiaria, FrmMarcacionesAutorizadas, FrmHorasTrabajadas, FrmSobreTiempo, FrmSuspenderAsistencia",
                "ClassRRHH.AsistenciaAccess; entidad Asistencia; spRRHH_GetAsistencias, spRRHH_GetDiasAsistidos, spRRHH_GetMinTarde, spRRHH_GetRecordAsistenciasSemanales",
            ],
            [
                "Importación de asistencia (múltiples fuentes)",
                "FrmImportarAsistencia, FrmImportarAsistenciaBio, FrmImportarAsistenciaCONSEJO, FrmImportarAsistenciaFueraSEDE, FrmImportarAsistenciaPROIND",
                "Inferido por nombres: integraciones/lote por sede o dispositivo biométrico",
            ],
            [
                "Permisos / salidas / faltas",
                "FrmPermiso, FrmMotivoPermiso, FrmMantTipoPerm, FrmMantMotivo, FrmConsultaPermisos, FrmConsultaFaltas, FrmSalida, FrmRelacionSalidaAutorizada, FrmFeriado",
                "spRRHH_llenarPermiso, spRRHH_getMotivo_perm, spRRHH_Faltas, spRRHH_RepPermisos",
            ],
        ],
    )

    doc.add_heading("1.4 Reportes", level=2)
    add_para(
        doc,
        "Múltiples formularios FrmReport* / FrmRep* y TableAdapters en DataSet1 "
        "consumen procedimientos de reporte (cuadro de personal, oficina, "
        "cumpleaños, cargos, profesión, fecha de ingreso, asistencia diaria, "
        "faltas, permisos, ficha). Evidencia: RecursosHumanos.xml y "
        "tipos spRRHH_Report_* / spRRHH_Rep_* en el exe y GRLL.",
    )
    add_para(
        doc,
        "Tecnología de reporting: DevExpress.XtraReports.v11.2 y "
        "Microsoft.ReportViewer.WinForms (presentes en el directorio raíz y Reportes/).",
    )

    doc.add_heading("1.5 Componentes de plataforma compartida (no exclusivos de RRHH)", level=2)
    add_para(
        doc,
        "GRLL.dll contiene además contratos a módulos ajenos al foco RRHH "
        "(p.ej. sp_aba_*, sp_Tra_*, spCJ_*, spCJR_*, spKM_*), lo que indica "
        "una biblioteca corporativa GRLL reutilizada. En este cliente solo "
        "se confirma uso intensivo del subconjunto RRHH/SEG vía formularios "
        "del exe; el uso runtime de ABA/Trámite/Caja/etc. desde este EXE "
        "no está determinado con certeza.",
    )

    # ========== 2. COMPONENTES CRÍTICOS ==========
    doc.add_heading("2. Componentes críticos y dependencias con la base de datos", level=1)
    add_para(
        doc,
        "Se listan puntos de los que dependen múltiples procesos. "
        "Donde el esquema físico no está en el repo, la dependencia de BD "
        "se expresa vía SP/entidades evidenciadas en el ensamblado.",
    )

    add_table(
        doc,
        ["Componente", "Qué hace", "Dependencias BD / servicio", "Qué se rompe si falla", "Evidencia"],
        [
            [
                "SQL Server / BD CMI",
                "Almacén central de datos de RRHH y seguridad de app",
                "Catálogo CMI; SP spRRHH_*, spSEG_*; entidades Empleado, Asistencia, Permiso, Usuario, Rol, etc.",
                "Toda la aplicación: login, fichas, asistencia, reportes",
                "RecursosHumanos.exe.config (Initial Catalog=CMI); connectionStrings; MySettings.BaseDatos",
            ],
            [
                "GRLL.Common.clsConexion",
                "Centraliza apertura/cierre de SqlConnection, transacciones y armado de connection string",
                "SQL Server vía System.Data.SqlClient; parámetros Servidor/BaseDatos/Usuario/Password/AutenticacionIntegrada",
                "Login y cualquier acceso a datos que use SQL_AbrirConexion / Conexion estática",
                "_auth_analysis/dec_GRLL_Common_clsConexion.cs (SQL_AbrirConexion, CrearCadenaConexion)",
            ],
            [
                "GRLL.Common.ucLogin + FrmLogin",
                "UI de autenticación; asigna Usuario/Password y prueba la conexión",
                "Autenticación SQL (no tabla de usuarios en el click); luego posible GRLL_Usuario_Acceso",
                "No se puede ingresar al sistema",
                "dec_GRLL_Common_ucLogin.cs btnAceptar_Click L426-458; exe FrmLogin; strings 'Acceso al Sistema'",
            ],
            [
                "SQLConfiguracion / MySettings",
                "Parámetros de servidor, BD, timeout, flag Integrated Security",
                "Indirecta: define a qué instancia/BD se conecta",
                "Conexiones fallidas o a instancia incorrecta",
                "exe.config applicationSettings; clsConexion.Configuracion",
            ],
            [
                "DataSet1 + TableAdapters spRRHH_*",
                "Capa tipada ADO.NET para fichas y reportes",
                "Procedimientos almacenados spRRHH_Ficha_*, spRRHH_Report_*, etc.; CMIConnectionString (sa) para diseñador/adapters",
                "Reportes y pantallas que llenan DataTables tipados",
                "RecursosHumanos.xml; tipos DataSet1* en exe_classes.txt",
            ],
            [
                "ClassRRHH (AsistenciaAccess, RRHHClass, FichaSocialClass, FamiliaresClass)",
                "Lógica de dominio RRHH empaquetada",
                "No determinado sin descompilar métodos; se asume consumo de BD vía GRLL/SqlClient",
                "Asistencia, ficha social y operaciones encapsuladas en esta DLL",
                "classrrhh_classes.txt; ClassRRHH.dll",
            ],
            [
                "Entidades Usuario / Rol / Rol_Acceso + spSEG_*",
                "Modelo de seguridad de aplicación (roles/interfaces)",
                "SP: spSEG_ListarUsuarios, spSEG_VerfUsuario, spSEG_llenarInterfaz*, spSEG_listarUserRol; entidades Usuario, Usuario_Rol, Rol, Rol_Acceso",
                "Autorización de menús/pantallas (post-login); detalle exacto de enforcement no trazado línea a línea",
                "grll_classes.txt; grll_login_types.txt",
            ],
            [
                "Entidad Empleado (+ Empleado_Area, Empleado_Marcacion)",
                "Núcleo de personal",
                "SP spEmpleado_*, spRRHH_busca_empleado, sp_RRHH_Busca_Empleado_tree",
                "Fichas, organigrama, horarios, reportes de personal",
                "grll_classes.txt Empleado / EmpleadoMapper",
            ],
            [
                "DevExpress / ReportViewer",
                "Controles de UI y reportes",
                "No son BD; dependen de ensamblados locales v11.2 (y legacy v7.3/v9.1)",
                "Pantallas/reportes no renderizan si faltan DLL",
                "root_binaries.txt",
            ],
        ],
    )

    # ========== 3. REGLAS DE NEGOCIO ==========
    doc.add_heading("3. Reglas de negocio identificadas", level=1)
    add_para(
        doc,
        "Cada regla se etiqueta como Confirmada (con evidencia concreta) o "
        "Inferida (evidencia indirecta). No se inventan constraints de BD "
        "ausentes del repositorio.",
    )

    add_table(
        doc,
        ["Regla", "Estado", "Evidencia"],
        [
            [
                "El login de la aplicación usa autenticación SQL Server con el usuario y contraseña capturados en el formulario (salvo AutenticacionIntegrada=True).",
                "Confirmada",
                "ucLogin.btnAceptar_Click asigna Configuracion.Usuario/Password; clsConexion.CrearCadenaConexion agrega User Id/Password; AutenticacionIntegrada=False en exe.config",
            ],
            [
                "Máximo 3 intentos fallidos de conexión en el control de login antes de forzar salida de intentos.",
                "Confirmada",
                "Contador estático en btnAceptar_Click; cadena 'Nº intentos = 3'",
            ],
            [
                "Existe un chequeo de 'acceso al sistema' a nivel aplicación distinto del login SQL.",
                "Confirmada (existencia del mensaje/tipos); flujo completo Inferido",
                "Cadena 'No tiene Acceso al Sistema'; tipos GRLL_Usuario_Acceso*; implementación exacta en FrmLogin no completamente descompilada aquí",
            ],
            [
                "La base de datos operativa se llama CMI.",
                "Confirmada",
                "exe.config Initial Catalog=CMI y MySettings.BaseDatos=CMI",
            ],
            [
                "Antes de eliminar documentos/periodos/rotaciones/asegurados se exige selección previa y/o confirmación.",
                "Confirmada (mensajes UI)",
                "Strings: 'Primero debe Seleccionar...', 'Desea eliminar el Documento/Periodo/Rotación...'",
            ],
            [
                "Validación de tamaño en tipo de sangre / datos maestros.",
                "Confirmada (mensaje)",
                "Cadena 'Tipo de Sangre: Tamaño Excedido' asociada a FrmMantGrupoSanguineo",
            ],
            [
                "Existe verificación de número de ley en régimen de pensiones.",
                "Confirmada (mensaje)",
                "'Verificar Nº de Ley' / FrmMantRegimen",
            ],
            [
                "Los reportes de ficha y personal se obtienen vía procedimientos almacenados dedicados, no SQL ad-hoc visible en el cliente.",
                "Confirmada (contratos)",
                "DataSet1 TableAdapters nombrados igual que SP spRRHH_*",
            ],
            [
                "El modelo de seguridad de aplicación se basa en Usuario–Rol–Rol_Acceso/Interfaces.",
                "Inferida",
                "Entidades y SP spSEG_* / mappers; sin DDL ni código de enforcement descompilado completo",
            ],
            [
                "Las importaciones de asistencia (Bio, CONSEJO, FueraSEDE, PROIND) representan reglas distintas por origen de marcación.",
                "Inferida",
                "Solo nombres de formularios FrmImportarAsistencia*; lógica interna no determinada",
            ],
            [
                "AnioActual=2013 en userSettings puede condicionar filtros/procesos por año.",
                "Inferida",
                "Valor en exe.config userSettings; uso en código no verificado",
            ],
            [
                "Constraints FK/CHECK/triggers en tablas Empleado/Asistencia/Permiso.",
                "No determinada",
                "No hay scripts SQL ni introspección de BD en el repositorio",
            ],
        ],
    )

    # ========== 4. ÁREAS SIN DOCUMENTACIÓN ==========
    doc.add_heading("4. Áreas con poca o ninguna documentación", level=1)

    doc.add_heading("4.1 Ausencia de documentación de producto", level=2)
    bullets_doc = [
        "No existe README, wiki, manual de usuario ni especificación funcional en el repositorio.",
        "Los XML de documentación de ensamblado (ClassRRHH.xml, gran parte de RecursosHumanos.xml) contienen principalmente comentarios auto-generados de recursos/My Project, no reglas de negocio.",
        "Las rutas PDB apuntan a máquinas de desarrollo históricas (D:\\EVV\\..., E:\\EVV\\Proyectos GRLL\\...) no presentes aquí.",
    ]
    for b in bullets_doc:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("4.2 Esquema de base de datos — no determinado", level=2)
    bullets_bd = [
        "No hay archivos .sql, diagramas ni export de tablas/vistas/triggers/índices.",
        "No se pueden afirmar PK/FK reales; solo nombres de entidades LINQ/mappers y SP.",
        "Cuerpo (T-SQL) de spRRHH_* y spSEG_* no está en el repo; solo firmas/result types en GRLL.",
        "Relación exacta entre login SQL Server y filas de tabla Usuario / spSEG_VerfUsuario: no determinada.",
    ]
    for b in bullets_bd:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("4.3 Código fuente y lógica opaca", level=2)
    bullets_code = [
        "No hay proyectos .vbproj/.sln ni fuentes .vb en el workspace; solo binarios + PDB + descompilados parciales en _auth_analysis.",
        "FrmLogin.vb (origen PDB) no está descompilado línea a línea en este paquete: el cableado exacto Login_onAceptar → GRLL_Usuario_Acceso queda parcialmente no determinado.",
        "Métodos internos de ClassRRHH.AsistenciaAccess / RRHHClass / FichaSocialClass: no documentados (ClassRRHH.xml vacío de API).",
        "ClsCtrls.dll y GRLL.Common.Rutinas.dll: rol detallado no analizado en profundidad.",
        "No hay pruebas automatizadas (ni proyectos de test) visibles.",
        "Configuración RutaLogo apunta a D:\\Sistema_GRLL\\Logo_GRLL.jpg — ruta de entorno original; impacto en runtime local no verificado.",
        "Uso real desde este EXE de módulos GRLL no-RRHH (ABA, Trámite, Caja, CJR, KM): no determinado.",
    ]
    for b in bullets_code:
        doc.add_paragraph(b, style="List Bullet")

    # ========== 5. LIMITACIONES LEGADO ==========
    doc.add_heading("5. Limitaciones del sistema legado", level=1)
    add_table(
        doc,
        ["Limitación", "Evidencia / patrón"],
        [
            [
                "Arquitectura monolítica de escritorio sin capa API/servicios HTTP",
                "Cliente WinForms (RecursosHumanos.exe) habla directo a SQL Server vía ADO.NET/SqlClient",
            ],
            [
                "Autenticación acoplada a logins de SQL Server",
                "ucLogin + CrearCadenaConexion con User Id del formulario; no Identity/AD corporativo como mecanismo principal (AutenticacionIntegrada=False)",
            ],
            [
                "Credenciales en texto plano en configuración",
                "connectionStrings con User ID=sa y Password en RecursosHumanos.exe.config",
            ],
            [
                "Dos modelos de conexión potencialmente divergentes",
                "CMIConnectionString (sa) para DataSets vs. cadena dinámica de clsConexion con usuario de login",
            ],
            [
                "Biblioteca corporativa sobrecargada (GRLL)",
                "GRLL.dll ~2.2 MB con cientos de SP de múltiples sistemas; alto acoplamiento y superficie de cambio",
            ],
            [
                "Dependencias UI obsoletas / multi-versión",
                "Conviven DevExpress v7.3, v9.1 y v11.2 en el mismo despliegue",
            ],
            [
                "Criptografía débil / obsoleta en Seguridad.dll",
                "TripleDES + MD5, modo ECB por defecto (dec_Seguridad.cs EncryptString/DecryptString)",
            ],
            [
                "Manejo de errores genérico en login",
                "catch Exception y MessageBox 'ERROR DE CONEXION' + Message de SqlException; rethrow como Exception nueva en clsConexion",
            ],
            [
                "Estado de conexión estático global",
                "clsConexion._Cn estático compartido — riesgo de concurrencia/fugas en escenarios multi-hilo",
            ],
            [
                "Ausencia de código fuente y de tests en el repositorio entregado",
                "Solo binarios/PDB/config; imposibilidad de mantenimiento seguro sin descompilar o recuperar fuentes",
            ],
            [
                "Configuración de entorno hardcodeada / desactualizada",
                "AnioActual=2013; RutaLogo a ruta absoluta de otra máquina; copias temp/ y Reportes/ con configs distintas (históricamente server1 / sa)",
            ],
            [
                "Escalabilidad limitada",
                "Cliente grueso + SP en SQL; sin evidencia de colas, caché distribuida ni separación lectura/escritura",
            ],
            [
                "Interop Office legado",
                "Interop.Excel.dll / Office.Core presentes — dependencia de Excel instalado para algunas rutinas",
            ],
        ],
    )

    # ========== ANEXO ==========
    doc.add_heading("Anexo A — Stack tecnológico observado", level=1)
    add_table(
        doc,
        ["Capa", "Tecnología"],
        [
            ["Lenguaje / runtime", "VB.NET / .NET Framework (My.* , System 2.x settings; WinForms)"],
            ["UI", "Windows Forms + DevExpress Xtra* v11.2 (legacy v7.3/v9.1)"],
            ["Datos", "ADO.NET SqlClient; Typed DataSets/TableAdapters; System.Data.Linq (mappers)"],
            ["BD", "Microsoft SQL Server — catálogo CMI (evidencia en config)"],
            ["Reporting", "DevExpress XtraReports; Microsoft ReportViewer WinForms"],
            ["Seguridad auxiliar", "Seguridad.dll — TripleDES/MD5"],
            ["Arquitectura", "Cliente monolítico + biblioteca compartida GRLL + DLL de dominio ClassRRHH"],
        ],
    )

    doc.add_heading("Anexo B — Estructura del repositorio (vista analizada)", level=1)
    for line in [
        "Raíz: RecursosHumanos.exe (+ .config, .pdb, .xml), GRLL.dll, ClassRRHH.dll, Seguridad.dll, ClsCtrls.dll, DevExpress*, ReportViewer*, Interop.*",
        "Reportes/: copia de despliegue orientada a reportes/localización (es/de/ja/ru)",
        "temp/: copia alternativa de binarios/config (posiblemente build o entorno previo)",
        "_auth_analysis/: artefactos de ingeniería inversa (descompilados, inventarios de tipos, strings)",
        "Ausente: solución Visual Studio, fuentes .vb, scripts SQL, README, tests",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Anexo C — Procedimientos almacenados RRHH/SEG evidentes (lista no exhaustiva)", level=1)
    add_para(
        doc,
        "Extraídos de nombres de tipos *Result / TableAdapters en GRLL y RecursosHumanos. "
        "No implica que todos se invoquen en cada sesión.",
    )
    sps = (
        "spRRHH_busca_empleado, spRRHH_BuscarEmpleado_Codigo, sp_RRHH_Busca_Empleado_tree, "
        "spRRHH_Ficha_DatosPersonales, spRRHH_Ficha_Det_Familiar, spRRHH_Ficha_Det_Familiar_Salud, "
        "spRRHH_Ficha_Estudio_Realizado, spRRHH_Ficha_Estudio_Titulo, spRRHH_Ficha_Estudio_Capacitacion, "
        "spRRHH_Ficha_ExpLab, spRRHH_Ficha_ExpLab_Datos, spRRHH_Ficha_Ingreso_Egreso, spRRHH_Ficha_Vivienda, "
        "spRRHH_GetAsistencias, spRRHH_GetDiasAsistidos, spRRHH_GetMinTarde, spRRHH_GetRecordAsistenciasSemanales, "
        "spRRHH_GetPeriodoLaboralByEmpleado, spRRHH_Faltas, spRRHH_llenarPermiso, spRRHH_getMotivo_perm, "
        "spRRHH_RepAsisdiarria, spRRHH_RepPermisos, spRRHH_Report_CuadroPer, spRRHH_Report_Oficina, "
        "spRRHH_Report_Cumple, spRRHH_Rep_cargos, spRRHH_Report_FechaIngreso, spRRHH_Rep_profesion, "
        "spRRHH_busca_Disciplina, spRRHH_busca_archivos, spRRHH_ListarRotaciones, spRRHH_ListarResolucion, "
        "spSEG_ListarUsuarios, spSEG_llenarUsuarios, spSEG_VerfUsuario, spSEG_llenarInterfaz, "
        "spSEG_llenarInterfaz_Rol, spSEG_listarUserRol, GRLL_Usuario_Acceso"
    )
    add_para(doc, sps)

    doc.add_paragraph()
    fin = doc.add_paragraph()
    fr = fin.add_run(
        "Fin del informe. Toda afirmación técnica anterior está anclada a "
        "artefactos del repositorio o a descompilaciones citadas. "
        "Lo no verificable se declara explícitamente como no determinado o inferido."
    )
    fr.italic = True
    fr.font.size = Pt(10)

    doc.save(str(OUT))
    print(f"OK:{OUT}")
    print(f"SIZE:{OUT.stat().st_size}")


if __name__ == "__main__":
    main()
