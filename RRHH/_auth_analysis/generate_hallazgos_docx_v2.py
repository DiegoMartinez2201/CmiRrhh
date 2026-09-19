# -*- coding: utf-8 -*-
"""Regenera Hallazgos_Analisis_Sistema.docx con evidencia de app + BD CMI."""
from pathlib import Path
import sys

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(r"D:\UPN\CICLO 10\Evo conf software\RRHH\RRHH\Hallazgos_Analisis_Sistema.docx")


def add_para(doc, text, bold=False, size=11, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Hallazgos de Análisis Funcional y Técnico", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Sistema de Recursos Humanos (RecursosHumanos / GRLL)\n"
        "Informe de ingeniería inversa — aplicación de escritorio + base de datos CMI"
    )
    r.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        "App: D:\\UPN\\CICLO 10\\Evo conf software\\RRHH\\RRHH\n"
        "BD: CMI (backup CMI_Backup_QA.bak + instancia DESKTOP-HNG11PV\\SQLEXPRESS)\n"
        "Fecha del informe: 10 de septiembre de 2026"
    )
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ========== RESUMEN EJECUTIVO ==========
    doc.add_heading("Resumen ejecutivo", level=1)
    for line in [
        "El repositorio de la aplicación es un despliegue compilado VB.NET WinForms (RecursosHumanos.exe + GRLL.dll + ClassRRHH.dll), sin fuentes .vb ni solución Visual Studio.",
        "La base operativa es SQL Server, catálogo CMI: backup completo CMI_Backup_QA.bak (~190 MB) y instancia local ONLINE en DESKTOP-HNG11PV\\SQLEXPRESS.",
        "En CMI se confirmaron 101 tablas de usuario, 303 procedimientos, 90 vistas, 11 triggers, 85 FK y 98 PK (conteos de catálogo).",
        "El login autentica como login de SQL Server con el usuario/clave del formulario; luego existe chequeo de aplicación vía GRLL_Usuario_Acceso (cuerpo cifrado).",
        "El núcleo funcional cubre escalafón/ficha, asistencia/marcación, permisos, maestros y reportes, apoyado en tablas Empleado/Asistencia/Permiso/Usuario y SP spRRHH_*/spSEG_*.",
        "Modelo de seguridad de app: Usuario → Usuario_Rol → Rol → Rol_Acceso → SistemaOpcion; Usuario.IdEmpleado → Empleado (FK confirmada).",
        "Asistencia y Permiso tienen triggers de auditoría (INSERT/UPDATE); volúmenes aproximados: Empleado 3.335, Asistencia 155.684, Permiso 26.120, Usuario 95.",
        "Puntos únicos de falla: SQL Server/CMI, clsConexion y el flujo de login; sin ellos no opera el sistema.",
        "No hay README/manual; varios SP de seguridad están cifrados; no hay CHECK constraints en Empleado/Asistencia/Permiso/Usuario.",
        "Deuda técnica: monolito de escritorio, credenciales en texto plano en .config, DevExpress multi-versión, TripleDES+MD5, sin API ni pruebas automatizadas.",
    ]:
        p = doc.add_paragraph(line, style="List Number")
        for run in p.runs:
            run.font.size = Pt(10)

    # ========== METODOLOGÍA ==========
    doc.add_heading("Metodología", level=1)
    add_para(
        doc,
        "Análisis estático de la aplicación y lectura de catálogo de la BD CMI. "
        "No se modificó esquema ni se restauró el .bak en esta corrida (la BD CMI ya estaba ONLINE).",
    )
    add_para(doc, "Fuentes revisadas:", bold=True)
    for f in [
        "App: RecursosHumanos.exe/.config/.pdb/.xml, GRLL.dll, ClassRRHH.dll, Seguridad.dll, ClsCtrls.dll, DevExpress*, ReportViewer*, Interop.*",
        "Carpetas app: raíz, Reportes/, temp/, _auth_analysis/ (descompilados ucLogin, clsConexion, clsLogin, Seguridad; inventarios de tipos; strings).",
        "BD backup: D:\\UPN\\CICLO 10\\Evo conf software\\CMI_Backup_QA\\CMI_Backup_QA.bak (RESTORE HEADERONLY/FILELISTONLY).",
        "BD live: sqlcmd -E sobre DESKTOP-HNG11PV\\SQLEXPRESS / CMI (sys.tables, sys.procedures, sys.foreign_keys, sys.triggers, OBJECT_DEFINITION, COUNT(*)).",
        "Búsqueda de README/DDL scripts en repo app: no encontrados.",
    ]:
        doc.add_paragraph(f, style="List Bullet")
    add_para(
        doc,
        "Limitación: cuerpos de GRLL_Usuario_Acceso, spSEG_VerfUsuario y spSEG_CrearUserSql están cifrados (is_encrypted=1); su lógica interna se declara no determinada.",
        italic=True,
    )

    # ========== 1. FUNCIONALIDADES ==========
    doc.add_heading("1. Funcionalidades principales", level=1)
    add_para(
        doc,
        "Identificadas por formularios Frm* del exe, strings de UI, TableAdapters DataSet1, "
        "entidades/SP en GRLL y objetos reales en BD CMI.",
    )

    doc.add_heading("1.1 Acceso y seguridad", level=2)
    add_table(
        doc,
        ["Funcionalidad", "Evidencia", "Descripción"],
        [
            [
                "Login / Acceso al Sistema",
                "FrmLogin; ucLogin; título UI 'Acceso al Sistema'",
                "Captura usuario/clave; btnAceptar_Click abre SqlConnection vía clsConexion.",
            ],
            [
                "Límite de 3 intentos",
                "ucLogin contador estático; cadena 'Nº intentos = 3'",
                "Tras 3 fallos dispara onSalirIntentos.",
            ],
            [
                "Acceso por sistema/login de app",
                "GRLL_Usuario_Acceso(@IdSistema char(7), @login char(20)); mensaje 'No tiene Acceso al Sistema'",
                "SP existe y está cifrado; se invoca tras conectar (flujo FrmLogin parcialmente no descompilado).",
            ],
            [
                "Administración de usuarios/roles",
                "Tablas Usuario, Usuario_Rol, Rol, Rol_Acceso, SistemaOpcion; SP spSEG_*",
                "Modelo relacional confirmado en CMI; pantallas de admin no inventariadas como Frm* RRHH específicas.",
            ],
            [
                "Creación de login SQL",
                "spSEG_CrearUserSql(@usuario char(20)) cifrado",
                "Indica acoplamiento app↔login SQL Server.",
            ],
            [
                "Cifrado auxiliar",
                "Seguridad.Seguridad.EncryptString/DecryptString",
                "TripleDES + MD5; no sustituye el login SQL.",
            ],
        ],
    )

    doc.add_heading("1.2 Escalafón / ficha de personal", level=2)
    add_table(
        doc,
        ["Funcionalidad", "UI / clases", "Objetos BD"],
        [
            [
                "Mantenimiento de empleado / ficha",
                "FrmNuevaFicha, FrmMantEmpleado_Basico; ClassRRHH.RRHHClass",
                "Tabla Empleado (+ Persona, Cargo, etc.); spRRHH_GrabarEmpleado, spRRHH_ActualizaEmpleado, spRRHH_Ficha_*",
            ],
            [
                "Periodo laboral, rotaciones, resoluciones",
                "FrmPeriodoLaboral, FrmRotaciones, FrmResoluciones",
                "PeriodoLaboral, Rotacion, Resolucion; spRRHH_GetPeriodoLaboralByEmpleado, spRRHH_ListarRotaciones, spRRHH_ListarResolucion",
            ],
            [
                "Estudios, títulos, capacitaciones",
                "FrmEstudiosRealizados, FrmTitulos, FrmCapacitaciones",
                "EstudiosRealizado, TitulosEmpleado, Capacitacion; spRRHH_Ficha_Estudio_*, spRRHH_buscarEstudios",
            ],
            [
                "Experiencia laboral",
                "FrmExpLaboral",
                "ExpLaboral; spRRHH_Ficha_ExpLab*",
            ],
            [
                "Familiares / ficha social / vivienda",
                "FrmFamiliar, FrmFichaSocial; FichaSocialClass, FamiliaresClass",
                "Familiar, Vivienda, RRHH_Salud, RRHH_AspSocio, RRHH_DinamicaFamiliar, etc.; spRRHH_Ficha_Det_Familiar*, spRRHH_Ficha_Vivienda",
            ],
            [
                "Organigrama / personal por área",
                "FrmListaOrganica",
                "EstructOrganiz, Empleado_Area; sp_RRHH_Busca_Empleado_tree",
            ],
            [
                "Maestros (civil, sangre, nacionalidad, régimen, docs, cargos, condición laboral)",
                "FrmMantEstCivil, FrmMantGrupoSanguineo, FrmMantNacionalidades, FrmMantRegimen, FrmMantTipoDocIdent, FrmMantCargos, FrmMantTipoTrab",
                "EstadoCivil, TipoSangre, Nacionalidad, RegimenPension, TipoDocID, Cargo, TipoTrabajador",
            ],
            [
                "Asegurados / derechohabientes",
                "FrmRegistroAsegurado, FrmBuscarDerechohabiente",
                "RRHH_Asegurado",
            ],
        ],
    )

    doc.add_heading("1.3 Asistencia, horarios y permisos", level=2)
    add_table(
        doc,
        ["Funcionalidad", "UI / clases", "Objetos BD"],
        [
            [
                "Horarios y asignación",
                "FrmHorarios, FrmMantHorarios, FrmAsignarHorario*, FrmLocalAsignado",
                "Horario, HorarioTemporal, Local; spRRHH_llenarHorarioTemporal; trigger trg_update_HorarioTemporal",
            ],
            [
                "Registro / consulta de asistencia",
                "FrmAsistenciaManual, FrmConsultaAsistencias*, ClassRRHH.AsistenciaAccess",
                "Asistencia (PK Fecha+IdEmpleado); spRRHH_GetAsistencias, spRRHH_GetAsistenciaDia, spRRHH_GetMinTarde, etc.; triggers tr_aud_asis_norm*",
            ],
            [
                "Importación biométrica / por sede",
                "FrmImportarAsistencia, FrmImportarAsistenciaBio, CONSEJO, FueraSEDE, PROIND",
                "spRRHH_ImportarAsistenciaBio, spRRHH_ImportarAsistenciaBioFacial, SpRRHH_Asistencia_Import*, SpRRHH_Marcacion_Import; tablas Marcacion/marcaciones",
            ],
            [
                "Permisos / faltas / feriados",
                "FrmPermiso, FrmMotivoPermiso, FrmMantTipoPerm, FrmConsultaPermisos, FrmConsultaFaltas, FrmFeriado",
                "Permiso, Motivo_Perm, TipoPermiso, RRHH_Feriado; spRRHH_llenarPermiso, spRRHH_Faltas, spRRHH_RepPermisos; triggers tr_aud_permiso*",
            ],
        ],
    )

    doc.add_heading("1.4 Reportes", level=2)
    add_para(
        doc,
        "Formularios FrmReport*/FrmRep* y TableAdapters DataSet1 consumen SP de reporte "
        "(cuadro de personal, oficina, cumpleaños, cargos, profesión, fecha ingreso, "
        "asistencia diaria, faltas, permisos, ficha). Tecnología: DevExpress XtraReports v11.2 "
        "y Microsoft ReportViewer. Evidencia: RecursosHumanos.xml + lista de SP en CMI.",
    )

    doc.add_heading("1.5 Inventario cuantitativo de la BD CMI", level=2)
    add_table(
        doc,
        ["Métrica", "Valor", "Fuente"],
        [
            ["USER_TABLE", "101", "sys.tables (live CMI)"],
            ["SQL_STORED_PROCEDURE", "303", "sys.procedures"],
            ["VIEW", "90", "sys.views"],
            ["SQL_TRIGGER", "11", "sys.triggers"],
            ["FOREIGN_KEY_CONSTRAINT", "85", "sys.foreign_keys"],
            ["PRIMARY_KEY_CONSTRAINT", "98", "sys.key_constraints"],
            ["SQL_SCALAR_FUNCTION", "47", "sys.objects"],
            ["SP RRHH/SEG/GRLL relevantes", "73 (filtro de nombres)", "catálogo live"],
            ["Filas Usuario / Empleado / Asistencia / Permiso", "95 / 3.335 / 155.684 / 26.120", "COUNT(*)"],
            ["Backup CMI_Backup_QA.bak", "~190 MB, full, BD CMI, 2026-03-29", "RESTORE HEADERONLY"],
        ],
    )

    # ========== 2. COMPONENTES CRÍTICOS ==========
    doc.add_heading("2. Componentes críticos y dependencias con la base de datos", level=1)
    add_table(
        doc,
        ["Componente", "Qué hace", "Dependencias BD/servicios", "Qué se rompe si falla", "Evidencia"],
        [
            [
                "SQL Server / BD CMI",
                "Almacén central de RRHH y seguridad",
                "101 tablas; SP spRRHH_*/spSEG_*/GRLL_*; vistas vs_RRHH_*, vs_SEG_usuario",
                "Toda la aplicación",
                "exe.config Initial Catalog=CMI; CMI_Backup_QA.bak; instancia ONLINE",
            ],
            [
                "GRLL.Common.clsConexion",
                "Abre/cierra SqlConnection, transacciones, arma connection string",
                "SqlClient hacia CMI; SQLConfiguracion",
                "Login y todo acceso vía SQL_AbrirConexion",
                "dec_GRLL_Common_clsConexion.cs (CrearCadenaConexion, SQL_AbrirConexion)",
            ],
            [
                "ucLogin + FrmLogin",
                "UI de autenticación",
                "Login SQL + (post) GRLL_Usuario_Acceso",
                "No se ingresa al sistema",
                "ucLogin.btnAceptar_Click L426-458; FrmLogin",
            ],
            [
                "Tabla Empleado",
                "Núcleo de personal",
                "FK a Persona, Cargo, Horario, Local, EstructOrganiz, AFP, etc.; referenciada por Asistencia, Permiso, Usuario, Familiar, …",
                "Fichas, organigrama, asistencia, permisos, reportes",
                "sys.tables Empleado; 20+ FK hijas/padres listadas en catálogo",
            ],
            [
                "Tabla Asistencia + triggers auditoría",
                "Marcaciones diarias",
                "FK IdEmpleado→Empleado; tr_aud_asis_norm / _Insert",
                "Consultas, reportes e importaciones de asistencia",
                "CMI.Asistencia; 155.684 filas",
            ],
            [
                "Tabla Permiso + Motivo_Perm",
                "Gestión de permisos",
                "FK a Empleado y Motivo_Perm; triggers tr_aud_permiso*",
                "Módulo de permisos/reportes asociados",
                "CMI.Permiso; 26.120 filas",
            ],
            [
                "Usuario / Rol / Rol_Acceso / SistemaOpcion",
                "Autorización de aplicación",
                "Cadena FK Usuario→Empleado; Usuario_Rol; Rol_Acceso→SistemaOpcion; SP spSEG_*, GRLL_Usuario_Acceso",
                "Acceso a menús/opciones tras login SQL exitoso",
                "Esquema CMI + tipos en GRLL",
            ],
            [
                "DataSet1 / TableAdapters spRRHH_*",
                "Capa tipada ADO.NET para fichas/reportes",
                "SP de ficha/reporte + CMIConnectionString (sa)",
                "Pantallas/reportes tipados",
                "RecursosHumanos.xml; exe.config connectionStrings",
            ],
            [
                "ClassRRHH.dll",
                "Lógica de dominio (AsistenciaAccess, FichaSocialClass, …)",
                "Consume BD vía GRLL/SqlClient (cuerpo no documentado en XML)",
                "Operaciones encapsuladas de asistencia/ficha social",
                "classrrhh_classes.txt; ClassRRHH.dll",
            ],
        ],
    )

    doc.add_heading("2.1 Relaciones FK clave (confirmadas en CMI)", level=2)
    add_table(
        doc,
        ["FK", "Relación"],
        [
            ["FK__Usuario__IdEmple__5C2E5663", "Usuario.IdEmpleado → Empleado.IdEmpleado"],
            ["FK__Usuario_R__IdUsu__6C59D134", "Usuario_Rol.IdUsuario → Usuario.IdUsuario"],
            ["FK__Usuario_R__IdRol__6B65ACFB", "Usuario_Rol.IdRol → Rol.IdRol"],
            ["FK__Rol_Acces__IdRol__766D4B53", "Rol_Acceso.IdRol → Rol.IdRol"],
            ["FK__Rol_Acces__IdSis__7579271A", "Rol_Acceso.IdSistemaOpcion → SistemaOpcion.IdSistemaOpcion"],
            ["FK__Asistenci__IdEmp__6B7099F3", "Asistencia.IdEmpleado → Empleado.IdEmpleado"],
            ["FK__Permiso__IdEmple__5A9B13AC", "Permiso.IdEmpleado → Empleado.IdEmpleado"],
            ["FK__Permiso__idMotiv__59A6EF73", "Permiso.idMotivo → Motivo_Perm.idMotivo"],
            ["FK__Empleado__IdPers__585DC57F", "Empleado.IdPersona → Persona.IdPersona"],
            ["FK__Empleado__IdCarg__52A4EC29", "Empleado.IdCargo → Cargo.IdCargo"],
            ["FK_Empleado_aspnet_Users", "Empleado.Userid → aspnet_Users.UserId"],
        ],
    )

    # ========== 3. REGLAS DE NEGOCIO ==========
    doc.add_heading("3. Reglas de negocio identificadas", level=1)
    add_table(
        doc,
        ["Regla", "Estado", "Evidencia"],
        [
            [
                "El login de la app usa autenticación SQL Server con usuario/clave del formulario cuando AutenticacionIntegrada=False.",
                "Confirmada",
                "ucLogin.btnAceptar_Click; clsConexion.CrearCadenaConexion User Id/Password; exe.config AutenticacionIntegrada=False",
            ],
            [
                "Máximo 3 intentos fallidos de conexión en el control de login.",
                "Confirmada",
                "Contador en btnAceptar_Click; cadena 'Nº intentos = 3'",
            ],
            [
                "Tras conectar, el acceso funcional se valida por sistema+login vía GRLL_Usuario_Acceso.",
                "Confirmada (existencia/contrato); cuerpo Inferido/no legible",
                "SP en CMI con params @IdSistema,@login; cifrado; mensaje UI 'No tiene Acceso al Sistema'",
            ],
            [
                "Un Usuario de aplicación puede vincularse a un Empleado.",
                "Confirmada",
                "FK Usuario.IdEmpleado → Empleado.IdEmpleado; columna Login char(20)",
            ],
            [
                "La autorización de opciones es por Rol y SistemaOpcion (matriz Rol_Acceso.Permiso).",
                "Confirmada (estructura); enforcement UI Inferido",
                "Tablas Rol, Rol_Acceso, SistemaOpcion + FK; código de menú no trazado línea a línea",
            ],
            [
                "Asistencia se identifica por (Fecha, IdEmpleado).",
                "Confirmada",
                "PK de Asistencia en CMI",
            ],
            [
                "Permiso se identifica por (IdEmpleado, NPermiso) y referencia Motivo_Perm.",
                "Confirmada",
                "PK + FK idMotivo en CMI",
            ],
            [
                "Cambios en Asistencia y Permiso se auditan automáticamente.",
                "Confirmada",
                "Triggers tr_aud_asis_norm*, tr_aud_permiso* (INSERT/UPDATE) habilitados",
            ],
            [
                "Al actualizar HorarioTemporal se registra FechaRegistro = GETDATE().",
                "Confirmada",
                "Trigger trg_update_HorarioTemporal",
            ],
            [
                "Empleado pertenece a una estructura orgánica (Year + idAreaOrganiz).",
                "Confirmada",
                "FK Empleado.(Year,idAreaOrganiz) → EstructOrganiz",
            ],
            [
                "Existen validaciones UI de selección previa antes de eliminar documentos/periodos/rotaciones.",
                "Confirmada",
                "Strings del exe: 'Primero debe Seleccionar...', 'Desea eliminar...'",
            ],
            [
                "spSEG_CrearUserSql crea/sincroniza logins SQL a partir del usuario de app.",
                "Inferida",
                "Nombre del SP + parámetro @usuario; cuerpo cifrado — no se leyó el T-SQL",
            ],
            [
                "CHECK constraints de dominio en Empleado/Asistencia/Permiso/Usuario.",
                "No determinadas / ausentes",
                "Consulta de CHECK: 0 en esas tablas",
            ],
            [
                "Reglas internas exactas de importación Bio/CONSEJO/PROIND.",
                "Inferidas (por nombres de Frm/SP)",
                "Cuerpos SP y código de formularios no descompilados en este informe",
            ],
        ],
    )

    # ========== 4. ÁREAS SIN DOCUMENTACIÓN ==========
    doc.add_heading("4. Áreas con poca o ninguna documentación", level=1)

    doc.add_heading("4.1 Documentación de producto", level=2)
    for b in [
        "No hay README, wiki ni manual de usuario en el repositorio de la aplicación.",
        "XML de ensamblado (ClassRRHH.xml, etc.) son stubs de diseñador/recursos, no reglas de negocio.",
        "PDB apunta a rutas históricas (D:\\EVV\\..., E:\\EVV\\Proyectos GRLL\\...) ausentes aquí.",
    ]:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("4.2 Base de datos — parcial / cifrado", level=2)
    for b in [
        "En CMI_Backup_QA solo hay .bak; no hay scripts DDL versionados en carpeta.",
        "GRLL_Usuario_Acceso, spSEG_VerfUsuario y spSEG_CrearUserSql: cuerpos cifrados → lógica exacta no determinada.",
        "Detalle de tablas Asistencia_Aud / Permiso_Aud (destino de auditoría): estructura no expandida en este informe.",
        "Uso real de tablas aspnet_* (Membership) desde el cliente WinForms RRHH: no determinado (existe FK Empleado.Userid → aspnet_Users).",
        "90 vistas (p.ej. vs_RRHH_*, vs_SEG_usuario): listado parcial; definición completa no transcrita aquí.",
    ]:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("4.3 Código de aplicación", level=2)
    for b in [
        "Sin .vbproj/.sln/.vb en el workspace; solo binarios + PDB + descompilados parciales.",
        "FrmLogin: cableado exacto onAceptar → GRLL_Usuario_Acceso no descompilado línea a línea.",
        "Métodos de ClassRRHH.AsistenciaAccess / RRHHClass / FichaSocialClass: sin documentación de API.",
        "ClsCtrls.dll y GRLL.Common.Rutinas.dll: rol detallado no analizado en profundidad.",
        "Ausencia de pruebas automatizadas visibles.",
        "Módulos GRLL no-RRHH embebidos en la DLL (ABA, Trámite, Caja, CJR, KM): uso desde este EXE no determinado.",
    ]:
        doc.add_paragraph(b, style="List Bullet")

    # ========== 5. LIMITACIONES ==========
    doc.add_heading("5. Limitaciones del sistema legado", level=1)
    add_table(
        doc,
        ["Limitación", "Evidencia"],
        [
            [
                "Arquitectura monolítica de escritorio sin API HTTP",
                "RecursosHumanos.exe → SqlClient directo a CMI",
            ],
            [
                "Autenticación acoplada a logins de SQL Server",
                "ucLogin + CrearCadenaConexion; spSEG_CrearUserSql sugiere creación de logins SQL",
            ],
            [
                "Credenciales en texto plano en configuración",
                "RecursosHumanos.exe.config connectionStrings (User ID=sa y Password)",
            ],
            [
                "Doble modelo de conexión (sa en DataSet vs usuario de login en clsConexion)",
                "CMIConnectionString vs SQLConfiguracion dinámica",
            ],
            [
                "SP de seguridad cifrados — dificulta auditoría/mantenimiento",
                "is_encrypted=1 en GRLL_Usuario_Acceso, spSEG_VerfUsuario, spSEG_CrearUserSql",
            ],
            [
                "Sin CHECK constraints en tablas núcleo",
                "0 CHECK en Empleado/Asistencia/Permiso/Usuario",
            ],
            [
                "Criptografía débil en Seguridad.dll",
                "TripleDES + MD5, ECB por defecto (dec_Seguridad.cs)",
            ],
            [
                "Dependencias UI multi-versión / obsoletas",
                "DevExpress v7.3, v9.1 y v11.2 coexisten en el despliegue",
            ],
            [
                "Estado de conexión estático global",
                "clsConexion._Cn estático",
            ],
            [
                "Manejo de errores genérico en login",
                "catch Exception + MessageBox 'ERROR DE CONEXION'",
            ],
            [
                "Sin código fuente ni tests en el entregable",
                "Solo binarios/PDB/config",
            ],
            [
                "Configuración de entorno desactualizada / hardcodeada",
                "AnioActual=2013; RutaLogo=D:\\Sistema_GRLL\\Logo_GRLL.jpg",
            ],
            [
                "Volumen alto de asistencia en cliente grueso",
                "Asistencia ~155k filas; lógica en SP + WinForms",
            ],
            [
                "Interop Office legado",
                "Interop.Excel.dll / Office.Core en el directorio",
            ],
            [
                "BD mezcla dominio RRHH con Membership ASP.NET y objetos de otros módulos",
                "aspnet_* + tablas Limitacion/Ubigeo con triggers ERwin; SP no-RRHH en GRLL",
            ],
        ],
    )

    # ========== ANEXOS ==========
    doc.add_heading("Anexo A — Stack tecnológico", level=1)
    add_table(
        doc,
        ["Capa", "Tecnología"],
        [
            ["Lenguaje / runtime", "VB.NET / .NET Framework WinForms (My.Settings)"],
            ["UI", "Windows Forms + DevExpress Xtra* v11.2 (legacy v7.3/v9.1)"],
            ["Datos app", "ADO.NET SqlClient; Typed DataSets; System.Data.Linq (mappers GRLL)"],
            ["BD", "Microsoft SQL Server — CMI (compatibilidad catálogo nivel 100; backup con motor 16.x)"],
            ["Reporting", "DevExpress XtraReports; Microsoft ReportViewer"],
            ["Seguridad auxiliar", "Seguridad.dll TripleDES/MD5"],
            ["Arquitectura", "Cliente monolítico + GRLL compartida + ClassRRHH"],
        ],
    )

    doc.add_heading("Anexo B — Estructura revisada", level=1)
    for line in [
        "RRHH\\RRHH: RecursosHumanos.exe + DLLs de negocio/UI + configs",
        "RRHH\\RRHH\\Reportes y temp: copias de despliegue",
        "RRHH\\RRHH\\_auth_analysis: artefactos de ingeniería inversa previos",
        "CMI_Backup_QA\\CMI_Backup_QA.bak: backup full de CMI (~190 MB)",
        "Instancia live: DESKTOP-HNG11PV\\SQLEXPRESS database CMI ONLINE",
        "Ausente en app: solución VS, fuentes .vb, scripts SQL versionados, README, tests",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Anexo C — Tablas de usuario en CMI (101)", level=1)
    add_para(
        doc,
        "AFP, AguaViv, AlumbradoViv, asis_norm, Asistencia, Asistencia_Aud, aspnet_Applications, "
        "aspnet_Membership, aspnet_Paths, aspnet_PersonalizationAllUsers, aspnet_PersonalizationPerUser, "
        "aspnet_Profile, aspnet_Roles, aspnet_SchemaVersions, aspnet_Users, aspnet_UsersInRoles, "
        "aspnet_WebEvent_Events, AuxRRHH_ReporteAsistencia, Capacitacion, Cargo, Comportamiento, "
        "ConservacionViv, Dim_Date, Discapacidad, dtproperties, Empleado, Empleado_Area, EstadoCivil, "
        "EstadoEstudio, EstructOrganiz, EstudiosRealizado, ExcretasViv, ExpLaboral, Familia, Familiar, "
        "Horario, HorarioTemporal, Inconveniente, Institucion, Limitacion, Local, Marcacion, marcaciones, "
        "MaterialViv, Motivo_Perm, MotivoBaja, Nacionalidad, PeriodoLaboral, Permiso, Permiso_Aud, Persona, "
        "RegAsisDiario, RegimenPension, Reporte_Asistencias, Resolucion, ResponsableXUO, Rol, Rol_Acceso, "
        "Rotacion, RRHH_AcudeEnferm, RRHH_Asegurado, RRHH_AspSocio, RRHH_DinamicaFamiliar, RRHH_Feriado, "
        "RRHH_FuncFam, RRHH_Preg1, RRHH_Preg2, RRHH_Preg3, RRHH_RelacHermano, RRHH_RelacPareja, RRHH_RelacPH, "
        "RRHH_Salud, RRHH_TempFecha, RRHH_TipoFamilia, SistemaOpcion, sysdiagrams, sysLogEvent, TenenciaViv, "
        "TipoComportamiento, TipoDoc, TipoDocID, TipoEstudio, TipoFamiliar, TipoInstitucion, TipoMoneda, "
        "TipoMovimiento, TipoPermiso, TipoRecurso, TipoResolucion, TipoSangre, TipoTrabajador, TipoTransaccion, "
        "TipoViv, TitulosEmpleado, UbicacionViv, Ubigeo, UnidadMedida, Usuario, Usuario_Rol, ViveCon, Vivienda",
        size=9,
    )

    doc.add_heading("Anexo D — Triggers (11)", level=1)
    add_table(
        doc,
        ["Trigger", "Tabla", "Evento", "Propósito observado"],
        [
            ["tr_aud_asis_norm", "Asistencia", "UPDATE", "Auditoría de cambios de asistencia"],
            ["tr_aud_asis_norm_Insert", "Asistencia", "INSERT", "Auditoría de altas de asistencia"],
            ["tr_aud_permiso", "Permiso", "UPDATE", "Auditoría de cambios de permiso"],
            ["tr_aud_permiso_Insert", "Permiso", "INSERT", "Auditoría de altas de permiso"],
            ["trg_update_HorarioTemporal", "HorarioTemporal", "UPDATE", "Actualiza FechaRegistro=GETDATE()"],
            ["tD_Limitacion / tU_Limitacion", "Limitacion", "DELETE/UPDATE", "Patrón ERwin SET NULL hacia Recomendacion"],
            ["tD_TipoDocID / tU_TipoDocID", "TipoDocID", "DELETE/UPDATE", "ERwin enforce FK Persona_TipoDocID"],
            ["tD_Ubigeo / tU_Ubigeo", "Ubigeo", "DELETE/UPDATE", "ERwin enforce FK Meta_Ubigeo"],
        ],
    )

    doc.add_heading("Anexo E — Columnas núcleo Usuario / Asistencia / Permiso", level=1)
    add_para(doc, "Usuario (PK IdUsuario): IdUsuario, Login, Descripcion, Fecha, Estado, IdEmpleado.", size=10)
    add_para(
        doc,
        "Asistencia (PK Fecha+IdEmpleado): Hor_Ent, Hor_Sal, Alm_Sal, Alm_Ent, Flag_Ent, Flag_Sal, Estado, idHorario, Usuario, PC.",
        size=10,
    )
    add_para(
        doc,
        "Permiso (PK IdEmpleado+NPermiso): FechaInicio/Fin, Hora_Sal/Ret, Dia, Retorno, idMotivo, Lugar, Referencia, Autorizacion, Obs, AutorizacionRRHH, Autorizado.",
        size=10,
    )

    doc.add_paragraph()
    fin = doc.add_paragraph()
    fr = fin.add_run(
        "Fin del informe. Afirmaciones técnicas ancladas a binarios/config de la app, "
        "descompilaciones citadas y/o catálogo de la BD CMI. "
        "Lo no verificable se declara como no determinado, cifrado o inferido."
    )
    fr.italic = True
    fr.font.size = Pt(10)

    doc.save(str(OUT))
    print(f"OK:{OUT}")
    print(f"SIZE:{OUT.stat().st_size}")


if __name__ == "__main__":
    main()
